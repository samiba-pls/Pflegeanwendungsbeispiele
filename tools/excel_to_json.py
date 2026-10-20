#!/usr/bin/env python3
"""Liest die Excel-Tabelle und schreibt data.json fuer die Website.
Aufruf:  python tools/excel_to_json.py [excel-datei] [ausgabe.json]
"""
import json, sys
from pathlib import Path
import openpyxl

src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("KI_Pflege_Anwendungsbeispiele.xlsx")
dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("data.json")

wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Anwendungen"]
head = [str(c.value).strip() if c.value else "" for c in ws[1]]
col = {h: i for i, h in enumerate(head)}
need = ["ID", "Projekt / Anwendung", "Anbieter / Träger", "Ort", "Pflege-Setting", "Prozess-Typ",
        "KI-Typ", "Kategorie", "Umsetzungsstatus"]
missing = [h for h in need if h not in col]
if missing:
    sys.exit(f"Spalten fehlen in der Excel-Datei: {missing}")

def val(row, name):
    i = col.get(name)
    v = row[i] if i is not None and i < len(row) else None
    return "" if v is None else str(v).strip()

items, warnings = [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    name = val(row, "Projekt / Anwendung")
    if not name:
        continue  # leere Zeilen ueberspringen
    item = {
        "id": val(row, "ID"),
        "name": name,
        "traeger": val(row, "Anbieter / Träger"),
        "ort": val(row, "Ort"),
        "setting": val(row, "Pflege-Setting"),
        "prozess": val(row, "Prozess-Typ"),
        "ki": val(row, "KI-Typ"),
        "kategorie": val(row, "Kategorie"),
        "status": val(row, "Umsetzungsstatus"),
        "beschreibung": val(row, "Kurzbeschreibung / Use Case"),
        "link": val(row, "Link"),
    }
    leer = [k for k in ("ort", "setting", "prozess", "ki", "kategorie", "status") if not item[k]]
    if leer:
        warnings.append(f"ID {item['id']} ({name}): leer -> {', '.join(leer)}")
    items.append(item)

dst.write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"{len(items)} Eintraege -> {dst}")
for w in warnings:
    print("WARNUNG:", w)
